#!/usr/bin/env python3
"""Fail closed on format, anonymity, geometry, and placeholder defects in final DOCX files."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def document_text(doc: Document) -> str:
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in doc.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(parts)


def assert_page_geometry(doc: Document) -> None:
    for section in doc.sections:
        assert abs(section.page_width.inches - 8.5) < 0.01
        assert abs(section.page_height.inches - 11.0) < 0.01
        assert abs(section.top_margin.inches - 1.0) < 0.01
        assert abs(section.right_margin.inches - 1.0) < 0.01
        assert abs(section.bottom_margin.inches - 1.0) < 0.01
        assert abs(section.left_margin.inches - 1.0) < 0.01
        assert abs(section.header_distance.inches - 0.492) < 0.01
        assert abs(section.footer_distance.inches - 0.492) < 0.01


def assert_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    assert style.font.name == "Times New Roman"
    assert abs(style.font.size.pt - 12.0) < 0.01
    assert abs(float(style.paragraph_format.line_spacing) - 2.0) < 0.01


def assert_text_typography(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        if not paragraph.text.strip():
            continue
        spacing = paragraph.paragraph_format.line_spacing
        if spacing is None:
            spacing = paragraph.style.paragraph_format.line_spacing
        assert spacing is not None and abs(float(spacing) - 2.0) < 0.01, paragraph.text[:80]
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            assert run.font.name == "Times New Roman", run.text[:80]
            assert run.font.size is not None and abs(run.font.size.pt - 12.0) < 0.01, run.text[:80]


def assert_table_geometry(doc: Document) -> None:
    for index, table in enumerate(doc.tables):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        layout = tbl_pr.find(qn("w:tblLayout"))
        assert tbl_w is not None and tbl_w.get(qn("w:type")) == "dxa" and int(tbl_w.get(qn("w:w"))) == 9360, index
        assert tbl_ind is not None and int(tbl_ind.get(qn("w:w"))) == 0, index
        assert layout is not None and layout.get(qn("w:type")) == "fixed", index
        widths = [int(node.get(qn("w:w"))) for node in table._tbl.tblGrid.findall(qn("w:gridCol"))]
        assert sum(widths) == 9360, (index, widths)
        for row in table.rows:
            assert len(row.cells) == len(widths), index
            cell_widths = []
            for cell in row.cells:
                width = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                assert width is not None and width.get(qn("w:type")) == "dxa", index
                cell_widths.append(int(width.get(qn("w:w"))))
            assert cell_widths == widths, (index, cell_widths, widths)


def xml_audit(path: Path) -> dict[str, int]:
    with zipfile.ZipFile(path) as archive:
        xml_parts = [
            name
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        ]
        text = "\n".join(archive.read(name).decode("utf-8") for name in xml_parts)
        return {
            "anchor_objects": len(re.findall(r"<wp:anchor\b", text)),
            "inline_objects": len(re.findall(r"<wp:inline\b", text)),
            "rsid_attributes": len(re.findall(r"\sw:rsid\w*=", text)),
        }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--body", type=Path, required=True)
    parser.add_argument("--cover", type=Path, required=True)
    parser.add_argument("--author", required=True)
    parser.add_argument("--affiliation", required=True)
    parser.add_argument("--email", required=True)
    parser.add_argument("--output", type=Path, default=Path("artifacts/qa/nccsrj_docx_audit.json"))
    args = parser.parse_args()

    body = Document(args.body)
    cover = Document(args.cover)
    for doc in (body, cover):
        assert_page_geometry(doc)
        assert_normal_style(doc)
        assert_text_typography(doc)
        assert_table_geometry(doc)

    body_text = document_text(body)
    cover_text = document_text(cover)
    for value in (args.author, args.affiliation, args.email, "/Users/"):
        assert value.lower() not in body_text.lower(), f"Anonymous body contains {value!r}"
    for value in (args.author, args.affiliation, args.email):
        assert value in cover_text, f"Cover is missing {value!r}"

    placeholder_patterns = (r"\{\{", r"<YOUR", r"\bTBD\b", r"\[AFFILIATION", r"\bTO CONFIRM\b")
    for pattern in placeholder_patterns:
        assert not re.search(pattern, body_text, flags=re.IGNORECASE), pattern
        assert not re.search(pattern, cover_text, flags=re.IGNORECASE), pattern

    assert body.core_properties.author in (None, "")
    assert body.core_properties.last_modified_by in (None, "")
    assert len(body.tables) == 6, len(body.tables)
    assert len(body.inline_shapes) == 5, len(body.inline_shapes)
    assert "References" in body_text
    assert "Abstract" in body_text
    keyword_match = re.search(r"Keywords:\s*([^\n]+)", body_text)
    assert keyword_match
    keyword_count = len([value for value in keyword_match.group(1).split(";") if value.strip()])
    assert 5 <= keyword_count <= 7

    body_xml = xml_audit(args.body)
    cover_xml = xml_audit(args.cover)
    assert body_xml["anchor_objects"] == 0
    assert body_xml["inline_objects"] == 5
    assert body_xml["rsid_attributes"] == 0
    assert cover_xml["anchor_objects"] == 0

    report = {
        "status": "pass",
        "body": str(args.body),
        "cover": str(args.cover),
        "body_tables": len(body.tables),
        "body_inline_figures": len(body.inline_shapes),
        "keywords": keyword_count,
        "text_typography": "Times New Roman 12 pt, double-spaced for all text-bearing body/table paragraphs",
        "body_xml": body_xml,
        "cover_xml": cover_xml,
        "anonymous_body_core_author": body.core_properties.author or "",
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
