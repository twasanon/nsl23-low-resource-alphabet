#!/usr/bin/env python3
"""Build the anonymous NCCSRJ manuscript and separate author cover DOCX files."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Callable

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


MODEL_ORDER = ("single_frame_cnn", "cnn_bilstm", "cnn_transformer")
MODEL_LABELS = {
    "single_frame_cnn": "Single-frame CNN",
    "cnn_bilstm": "CNN–BiLSTM",
    "cnn_transformer": "CNN–Transformer",
}
MODEL_TOKENS = {
    "single_frame_cnn": "SINGLE",
    "cnn_bilstm": "BILSTM",
    "cnn_transformer": "TRANSFORMER",
}
STRUCTURAL_MARKERS = {
    "TABLE_PROTOCOL",
    "TABLE_MODELS",
    "TABLE_PRIMARY",
    "TABLE_FOLDS",
    "TABLE_CONDITIONS",
    "TABLE_RUNTIME",
    "FIGURE_STUDY_FLOW",
    "FIGURE_ARCHITECTURES",
    "FIGURE_PRIMARY",
    "FIGURE_TRAINING",
    "FIGURE_LIGHTING",
}
FONT = "Times New Roman"
CONTENT_WIDTH_DXA = 9360


def set_run_font(
    run: Any,
    *,
    size: float = 12,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str = "000000",
) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style: Any, size: float, *, bold: bool = False, italic: bool = False) -> None:
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)


def configure_document(doc: Document, tokens: dict[str, Any]) -> None:
    section = doc.sections[0]
    section.page_width = Inches(tokens["page"]["width_inches"])
    section.page_height = Inches(tokens["page"]["height_inches"])
    margins = tokens["page"]["margins_inches"]
    section.top_margin = Inches(margins["top"])
    section.right_margin = Inches(margins["right"])
    section.bottom_margin = Inches(margins["bottom"])
    section.left_margin = Inches(margins["left"])
    section.header_distance = Inches(tokens["page"]["header_inches"])
    section.footer_distance = Inches(tokens["page"]["footer_inches"])

    normal = doc.styles["Normal"]
    set_style_font(normal, tokens["typography"]["body_size_pt"])
    normal.paragraph_format.space_before = Pt(tokens["typography"]["body_before_pt"])
    normal.paragraph_format.space_after = Pt(tokens["typography"]["body_after_pt"])
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    heading1 = doc.styles["Heading 1"]
    set_style_font(heading1, 12, bold=True)
    heading1.paragraph_format.space_before = Pt(12)
    heading1.paragraph_format.space_after = Pt(0)
    heading1.paragraph_format.line_spacing = 2.0
    heading1.paragraph_format.keep_with_next = True

    heading2 = doc.styles["Heading 2"]
    set_style_font(heading2, 12, bold=True, italic=True)
    heading2.paragraph_format.space_before = Pt(8)
    heading2.paragraph_format.space_after = Pt(0)
    heading2.paragraph_format.line_spacing = 2.0
    heading2.paragraph_format.keep_with_next = True

    for name in ("Caption",):
        style = doc.styles[name]
        set_style_font(style, tokens["captions"]["size_pt"])
        style.paragraph_format.space_before = Pt(tokens["captions"]["before_pt"])
        style.paragraph_format.space_after = Pt(tokens["captions"]["after_pt"])
        style.paragraph_format.line_spacing = 2.0

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    set_run_font(run, size=tokens["page_furniture"]["footer_size_pt"])
    run._r.extend([begin, instruction, separate])
    placeholder = paragraph.add_run("1")
    set_run_font(placeholder, size=tokens["page_furniture"]["footer_size_pt"])
    end_run = paragraph.add_run()
    set_run_font(end_run, size=tokens["page_furniture"]["footer_size_pt"])
    end_run._r.append(end)


def set_core_properties(doc: Document, *, title: str, author: str = "", keywords: str = "") -> None:
    props = doc.core_properties
    props.title = title
    props.subject = "NCCSRJ research manuscript"
    props.author = author
    props.last_modified_by = author
    props.keywords = keywords
    props.comments = ""
    props.category = "Research article"


def add_inline_runs(paragraph: Any, text: str, size: float = 12) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            set_run_font(paragraph.add_run(text[position : match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size)
        position = match.end()
    if position < len(text):
        set_run_font(paragraph.add_run(text[position:]), size=size)


def add_body_paragraph(doc: Document, text: str, *, first_line_indent: bool = True) -> Any:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.widow_control = True
    if first_line_indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.3)
    add_inline_runs(paragraph, text)
    return paragraph


def add_caption(doc: Document, text: str, *, above: bool = False) -> Any:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_with_next = above
    add_inline_runs(paragraph, text, size=12)
    return paragraph


def add_table_text(doc: Document, text: str, *, above: bool) -> Any:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.keep_with_next = above
    add_inline_runs(paragraph, text, size=12)
    return paragraph


def set_cell_margins(cell: Any, *, top: int = 80, bottom: int = 80, start: int = 90, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell: Any, dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(dxa))
    width.set(qn("w:type"), "dxa")


def set_table_borders(table: Any, color: str = "666666", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def fixed_table(doc: Document, headers: list[str], rows: list[list[str]], widths_inches: list[float]) -> Any:
    if abs(sum(widths_inches) - 6.5) > 1e-6:
        raise ValueError(f"Table widths must sum to 6.5 inches: {widths_inches}")
    widths = [round(value * 1440) for value in widths_inches]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    set_table_borders(table)

    all_rows = [headers, *rows]
    for row_index, values in enumerate(all_rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(repeat)
        for column, (cell, value, width) in enumerate(zip(row.cells, values, widths, strict=True)):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F2F2F2")
                cell._tc.get_or_add_tcPr().append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 2.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(value))
            set_run_font(run, size=12, bold=row_index == 0)
    return table


def set_alt_text(picture_paragraph: Any, description: str) -> None:
    for doc_pr in picture_paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", description)
        doc_pr.set("title", description.split(".")[0])


def add_figure(doc: Document, path: Path, caption: str, alt_text: str, width: float = 6.3) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    set_alt_text(paragraph, alt_text)
    add_caption(doc, caption)


def percent(mean: float, sd: float | None = None) -> str:
    if sd is None:
        return f"{100.0 * mean:.1f}%"
    return f"{100.0 * mean:.1f} ± {100.0 * sd:.1f}%"


def number(value: float, decimals: int = 1) -> str:
    return f"{value:.{decimals}f}"


def load_results(results_dir: Path, runtime_path: Path) -> dict[str, pd.DataFrame]:
    paths = {
        "models": results_dir / "model_summary.csv",
        "folds": results_dir / "fold_summary.csv",
        "conditions": results_dir / "condition_summary.csv",
        "sign_types": results_dir / "sign_type_summary.csv",
        "paired": results_dir / "paired_differences_summary.csv",
        "runtime": runtime_path,
        "runs": results_dir / "run_metrics.csv",
        "histories": results_dir / "training_histories_all.csv",
    }
    missing = [str(path) for path in paths.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Manuscript result inputs are missing: {missing}")
    return {name: pd.read_csv(path) for name, path in paths.items()}


def load_audits(audit_dir: Path, config_path: Path) -> dict[str, Any]:
    paths = {
        "manifest": audit_dir / "nsl23_manifest_summary.json",
        "folds": audit_dir / "isolated_twofold_summary.json",
        "aggregate": audit_dir / "aggregate_exclusions_summary.json",
        "composition": audit_dir / "primary_set_composition.json",
        "aggregate_rows": audit_dir / "aggregate_exclusions.csv",
        "config": config_path,
    }
    missing = [str(path) for path in paths.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Manuscript audit inputs are missing: {missing}")
    return {
        "manifest": json.loads(paths["manifest"].read_text(encoding="utf-8")),
        "folds": json.loads(paths["folds"].read_text(encoding="utf-8")),
        "aggregate": json.loads(paths["aggregate"].read_text(encoding="utf-8")),
        "composition": json.loads(paths["composition"].read_text(encoding="utf-8")),
        "aggregate_rows": pd.read_csv(paths["aggregate_rows"]),
        "config": json.loads(paths["config"].read_text(encoding="utf-8")),
    }


def build_replacements(data: dict[str, pd.DataFrame], results_dir: Path) -> dict[str, str]:
    replacements: dict[str, str] = {}
    models = data["models"].set_index("model")
    metric_tokens = {
        "accuracy": "ACC",
        "macro_precision": "MACRO_P",
        "macro_recall": "MACRO_R",
        "macro_f1": "MACRO_F1",
        "weighted_f1": "WEIGHTED_F1",
    }
    for model in MODEL_ORDER:
        prefix = MODEL_TOKENS[model]
        row = models.loc[model]
        replacements[f"{prefix}_MODEL_LABEL"] = MODEL_LABELS[model]
        for metric, token in metric_tokens.items():
            replacements[f"{prefix}_{token}_PCT"] = f"{100 * row[f'{metric}_mean']:.1f}"
            replacements[f"{prefix}_{token}_SD_PCT"] = f"{100 * row[f'{metric}_sd']:.1f}"
            replacements[f"{prefix}_{token}_PM"] = percent(row[f"{metric}_mean"], row[f"{metric}_sd"])

    folds = data["folds"].set_index(["model", "fold"])
    for model in MODEL_ORDER:
        prefix = MODEL_TOKENS[model]
        for fold, fold_token in (("heldout_s1", "S1"), ("heldout_s2", "S2")):
            row = folds.loc[(model, fold)]
            replacements[f"{prefix}_{fold_token}_ACC_PM"] = percent(row["test_accuracy_mean"], row["test_accuracy_sd"])
            replacements[f"{prefix}_{fold_token}_F1_PM"] = percent(row["test_macro_f1_mean"], row["test_macro_f1_sd"])

    conditions = data["conditions"].copy()
    conditions["lighting"] = conditions["lighting"].str.lower()
    conditions = conditions.set_index(["model", "lighting"])
    for model in MODEL_ORDER:
        prefix = MODEL_TOKENS[model]
        for lighting in ("bright", "dark"):
            row = conditions.loc[(model, lighting)]
            light = lighting.upper()
            replacements[f"{prefix}_{light}_ACC_PM"] = percent(row["accuracy_mean"], row["accuracy_sd"])
            replacements[f"{prefix}_{light}_F1_PM"] = percent(row["macro_f1_mean"], row["macro_f1_sd"])

    sign_types = data["sign_types"].set_index(["model", "sign_type"])
    for model in MODEL_ORDER:
        prefix = MODEL_TOKENS[model]
        for sign_type in ("consonant", "vowel"):
            row = sign_types.loc[(model, sign_type)]
            token = sign_type.upper()
            replacements[f"{prefix}_{token}_ACC_PM"] = percent(row["accuracy_mean"], row["accuracy_sd"])
            replacements[f"{prefix}_{token}_F1_PM"] = percent(
                row["macro_f1_supported_classes_mean"], row["macro_f1_supported_classes_sd"]
            )

    paired = data["paired"].set_index(["comparison", "metric"])
    comparisons = {
        "TRANSFORMER_MINUS_SINGLE": "cnn_transformer_minus_single_frame_cnn",
        "TRANSFORMER_MINUS_BILSTM": "cnn_transformer_minus_cnn_bilstm",
        "BILSTM_MINUS_SINGLE": "cnn_bilstm_minus_single_frame_cnn",
    }
    for token, comparison in comparisons.items():
        for metric, metric_token in (("accuracy", "ACC"), ("macro_f1", "F1")):
            row = paired.loc[(comparison, metric)]
            replacements[f"{token}_{metric_token}_PP"] = f"{100 * row['difference_mean']:+.1f}"
            replacements[f"{token}_{metric_token}_SD_PP"] = f"{100 * row['difference_sd']:.1f}"

    runtime = data["runtime"].set_index("model")
    for model in MODEL_ORDER:
        prefix = MODEL_TOKENS[model]
        row = runtime.loc[model]
        replacements[f"{prefix}_MODEL_MEDIAN_MS"] = number(row["model_only_median_ms"])
        replacements[f"{prefix}_MODEL_P95_MS"] = number(row["model_only_p95_ms"])
        replacements[f"{prefix}_E2E_MEDIAN_MS"] = number(row["end_to_end_median_ms"])
        replacements[f"{prefix}_E2E_P95_MS"] = number(row["end_to_end_p95_ms"])
        replacements[f"{prefix}_E2E_CPS"] = number(row["end_to_end_clips_per_second"], 2)

    runs = data["runs"].sort_values(["model", "fold", "seed"]).drop_duplicates("model").set_index("model")
    for model in MODEL_ORDER:
        prefix = MODEL_TOKENS[model]
        replacements[f"{prefix}_TOTAL_PARAMS"] = f"{int(runs.loc[model, 'parameters_total']):,}"
        replacements[f"{prefix}_TRAINABLE_PARAMS"] = f"{int(runs.loc[model, 'parameters_trainable']):,}"

    final_history = data["histories"][data["histories"]["epoch"] == 30]
    for model in MODEL_ORDER:
        prefix = MODEL_TOKENS[model]
        rows = final_history[final_history["model"] == model]
        if len(rows) != 6:
            raise ValueError(f"Expected six final-epoch rows for {model}, found {len(rows)}")
        for column, token in (("train_accuracy", "FINAL_TRAIN_ACC"), ("train_macro_f1", "FINAL_TRAIN_F1")):
            replacements[f"{prefix}_{token}_PM"] = percent(rows[column].mean(), rows[column].std(ddof=1))
        replacements[f"{prefix}_FINAL_TRAIN_LOSS_PM"] = f"{rows['train_loss'].mean():.3f} ± {rows['train_loss'].std(ddof=1):.3f}"
        replacements[f"{prefix}_TRAIN_TEST_ACC_GAP_PP"] = f"{100 * (rows['train_accuracy'].mean() - models.loc[model, 'accuracy_mean']):.1f}"

    best_accuracy = models["accuracy_mean"].idxmax()
    best_f1 = models["macro_f1_mean"].idxmax()
    replacements["BEST_ACCURACY_MODEL"] = MODEL_LABELS[best_accuracy]
    replacements["BEST_MACRO_F1_MODEL"] = MODEL_LABELS[best_f1]
    replacements["BEST_ACCURACY_PM"] = percent(
        models.loc[best_accuracy, "accuracy_mean"], models.loc[best_accuracy, "accuracy_sd"]
    )
    replacements["BEST_MACRO_F1_PM"] = percent(
        models.loc[best_f1, "macro_f1_mean"], models.loc[best_f1, "macro_f1_sd"]
    )

    confusion_path = results_dir / "top_confusions_cnn_transformer.csv"
    confusions = pd.read_csv(confusion_path, keep_default_na=False).head(5)
    replacements["TRANSFORMER_TOP_CONFUSIONS"] = "; ".join(
        f"{row.true_label}→{row.predicted_label} ({int(row.count_across_seeds_and_folds)}/"
        f"{int(row.repeated_true_class_predictions)})"
        for row in confusions.itertuples()
    )
    return replacements


def replace_tokens(text: str, replacements: dict[str, str]) -> str:
    for key, value in replacements.items():
        text = text.replace("{{" + key + "}}", value)
    unresolved = sorted(
        token
        for token in set(re.findall(r"\{\{[A-Z0-9_]+\}\}", text))
        if token[2:-2] not in STRUCTURAL_MARKERS
    )
    if unresolved:
        raise ValueError(f"Unresolved manuscript tokens: {unresolved}")
    return text


def parse_blocks(source: str) -> tuple[str, str, str, list[tuple[str, str]]]:
    raw_blocks = [block.strip() for block in re.split(r"\n\s*\n", source) if block.strip()]
    if not raw_blocks or not raw_blocks[0].startswith("# "):
        raise ValueError("Manuscript source must begin with a level-one title")
    title = raw_blocks.pop(0)[2:].strip()
    if raw_blocks.pop(0) != "**Abstract**":
        raise ValueError("Title must be followed by an **Abstract** block")
    abstract = " ".join(raw_blocks.pop(0).splitlines())
    keyword_block = raw_blocks.pop(0)
    if not keyword_block.startswith("**Keywords:**"):
        raise ValueError("Abstract must be followed by a **Keywords:** block")
    keywords = keyword_block.removeprefix("**Keywords:**").strip()
    blocks: list[tuple[str, str]] = []
    for block in raw_blocks:
        if block.startswith("## "):
            blocks.append(("heading1", block[3:].strip()))
        elif block.startswith("### "):
            blocks.append(("heading2", block[4:].strip()))
        elif re.fullmatch(r"\{\{[A-Z0-9_]+\}\}", block):
            blocks.append(("marker", block[2:-2]))
        else:
            blocks.append(("paragraph", " ".join(block.splitlines())))
    return title, abstract, keywords, blocks


def manuscript_word_count(abstract: str, blocks: list[tuple[str, str]]) -> int:
    text_parts = [abstract]
    in_references = False
    for kind, text in blocks:
        if kind == "heading1" and text.lower() == "references":
            in_references = True
            continue
        if not in_references and kind == "paragraph":
            text_parts.append(re.sub(r"\*|`", "", text))
    return len(re.findall(r"\b[\w’'-]+\b", " ".join(text_parts), flags=re.UNICODE))


def table_protocol(doc: Document, audits: dict[str, Any]) -> None:
    manifest = audits["manifest"]
    folds = audits["folds"]
    aggregate = audits["aggregate"]
    composition = audits["composition"]
    config = audits["config"]
    overlap = int(audits["aggregate_rows"]["is_cropped"].astype(str).str.lower().isin({"true", "1", "yes"}).sum())
    fold1 = folds["folds"]["heldout_s1"]
    fold2 = folds["folds"]["heldout_s2"]
    run_count = len(config["models"]) * len(config["seed_list"]) * len(folds["folds"])
    add_table_text(doc, "**Table 1.** Audited dataset and evaluation protocol", above=True)
    rows = [
        ["Official release", f"NSL23 version 2; {manifest['decoded_videos']} decodable video files; {len(manifest['primary_usable_isolated_classes'])} alphabet classes"],
        ["Primary set", f"{manifest['primary_usable_isolated_clips']} original, individually labelled clips; {len(manifest['primary_usable_isolated_signers'])} signers"],
        ["Excluded sources", f"{aggregate['aggregate_sources']} aggregate recordings without timestamps; {manifest['cropped_or_preprocessed_videos']} cropped/preprocessed files ({overlap} overlap)"],
        ["Fold 1", f"Train {fold1['train_samples']} clips from {len(fold1['train_signers'])} signers; test {fold1['test_signer']} on {fold1['test_samples']} clips"],
        ["Fold 2", f"Train {fold2['train_samples']} clips from {len(fold2['train_signers'])} signers; test {fold2['test_signer']} on {fold2['test_samples']} clips"],
        ["Test conditions", f"{fold1['test_per_condition']['Indoor Unprepared - Bright']} bright and {fold1['test_per_condition']['Indoor Unprepared - Dark']} dark clips per held-out signer; every class in both conditions"],
        ["Vowel training", "Only S1 and S2 have original isolated vowels; each fold trains vowels from the other complete signer"],
        ["Selection", f"No validation set; fixed epoch {config['training']['max_epochs']} evaluated once per model/fold/seed"],
        ["Repetitions", f"Seeds {', '.join(map(str, config['seed_list']))}; {run_count} primary runs"],
    ]
    fixed_table(doc, ["Item", "Audited specification"], rows, [1.65, 4.85])
    add_table_text(doc, f"*Note.* Aggregate and cropped/preprocessed exclusions overlap by {overlap} files; the primary count uses the union of exclusions. The primary set contains {composition['sign_type']['consonant']} consonant and {composition['sign_type']['vowel']} vowel clips.", above=False)


def table_models(doc: Document, data: dict[str, pd.DataFrame], audits: dict[str, Any]) -> None:
    metrics = data["runs"].sort_values(["model", "fold", "seed"]).drop_duplicates("model").set_index("model")
    config = audits["config"]
    bilstm = config["models"]["cnn_bilstm"]
    transformer = config["models"]["cnn_transformer"]
    heads = {
        "single_frame_cnn": (f"Middle of {config['num_frames']}", "None; dropout + linear"),
        "cnn_bilstm": (f"{config['num_frames']} frames", f"{bilstm['projection_dim']}-D projection; {bilstm['layers']}-layer BiLSTM ({bilstm['hidden_dim']} per direction); layer norm; mean pool"),
        "cnn_transformer": (f"{config['num_frames']} frames", f"{transformer['projection_dim']}-D projection; {transformer['layers']} layers, {transformer['heads']} heads; final layer norm; mean pool"),
    }
    rows = []
    for model in MODEL_ORDER:
        frame_policy, head = heads[model]
        row = metrics.loc[model]
        rows.append(
            [
                MODEL_LABELS[model],
                frame_policy,
                head,
                f"{int(row['parameters_trainable']):,}",
                f"{int(row['parameters_total']):,}",
            ]
        )
    add_table_text(doc, "**Table 2.** Model comparison with a shared frozen MobileNetV2 frame encoder", above=True)
    fixed_table(doc, ["Model", "Input", "Temporal/classification head", "Trainable", "Total"], rows, [1.35, 0.9, 2.35, 0.9, 1.0])
    add_table_text(doc, "*Note.* Total parameters include the frozen ImageNet-pretrained MobileNetV2 encoder.", above=False)


def table_primary(doc: Document, data: dict[str, pd.DataFrame]) -> None:
    frame = data["models"].set_index("model")
    rows = []
    for model in MODEL_ORDER:
        row = frame.loc[model]
        rows.append(
            [
                MODEL_LABELS[model],
                percent(row["accuracy_mean"], row["accuracy_sd"]),
                percent(row["macro_precision_mean"], row["macro_precision_sd"]),
                percent(row["macro_recall_mean"], row["macro_recall_sd"]),
                percent(row["macro_f1_mean"], row["macro_f1_sd"]),
            ]
        )
    add_table_text(doc, "**Table 3.** Pooled out-of-fold performance across three fixed seeds", above=True)
    fixed_table(doc, ["Model", "Accuracy", "Macro precision", "Macro recall", "Macro-F1"], rows, [1.4, 1.15, 1.35, 1.3, 1.3])
    add_table_text(doc, "*Note.* Values are mean ± sample SD across three seeds after pooling both held-out folds (196 clips per seed). Metrics use the fixed 49-class order. Equal support for every held-out class makes macro recall numerically equal to accuracy.", above=False)


def table_folds(doc: Document, data: dict[str, pd.DataFrame]) -> None:
    frame = data["folds"].set_index(["model", "fold"])
    rows = []
    for model in MODEL_ORDER:
        s1 = frame.loc[(model, "heldout_s1")]
        s2 = frame.loc[(model, "heldout_s2")]
        rows.append(
            [
                MODEL_LABELS[model],
                percent(s1["test_accuracy_mean"], s1["test_accuracy_sd"]),
                percent(s1["test_macro_f1_mean"], s1["test_macro_f1_sd"]),
                percent(s2["test_accuracy_mean"], s2["test_accuracy_sd"]),
                percent(s2["test_macro_f1_mean"], s2["test_macro_f1_sd"]),
            ]
        )
    add_table_text(doc, "**Table 4.** Performance by held-out signer", above=True)
    fixed_table(doc, ["Model", "S1 accuracy", "S1 macro-F1", "S2 accuracy", "S2 macro-F1"], rows, [1.4, 1.25, 1.3, 1.25, 1.3])
    add_table_text(doc, "*Note.* Values are mean ± sample SD over seeds within each 98-clip signer fold.", above=False)


def table_conditions(doc: Document, data: dict[str, pd.DataFrame]) -> None:
    frame = data["conditions"].copy()
    frame["lighting"] = frame["lighting"].str.lower()
    frame = frame.set_index(["model", "lighting"])
    rows = []
    for model in MODEL_ORDER:
        bright = frame.loc[(model, "bright")]
        dark = frame.loc[(model, "dark")]
        rows.append(
            [
                MODEL_LABELS[model],
                percent(bright["accuracy_mean"], bright["accuracy_sd"]),
                percent(dark["accuracy_mean"], dark["accuracy_sd"]),
                percent(bright["macro_f1_mean"], bright["macro_f1_sd"]),
                percent(dark["macro_f1_mean"], dark["macro_f1_sd"]),
            ]
        )
    add_table_text(doc, "**Table 5.** Descriptive bright- and dark-condition performance", above=True)
    fixed_table(doc, ["Model", "Bright accuracy", "Dark accuracy", "Bright macro-F1", "Dark macro-F1"], rows, [1.4, 1.25, 1.25, 1.3, 1.3])
    add_table_text(doc, "*Note.* Values are mean ± sample SD over seeds. Each condition contains 98 pooled out-of-fold clips per seed. Recording condition was not randomized, so differences are descriptive.", above=False)


def table_runtime(doc: Document, data: dict[str, pd.DataFrame]) -> None:
    frame = data["runtime"].set_index("model")
    rows = []
    for model in MODEL_ORDER:
        row = frame.loc[model]
        rows.append(
            [
                MODEL_LABELS[model],
                f"{int(row['parameters_trainable']):,}",
                f"{row['model_only_median_ms']:.1f} ({row['model_only_p95_ms']:.1f})",
                f"{row['end_to_end_median_ms']:.1f} ({row['end_to_end_p95_ms']:.1f})",
                f"{row['end_to_end_clips_per_second']:.2f}",
            ]
        )
    add_table_text(doc, "**Table 6.** Synchronized batch-one Apple M4 MPS latency", above=True)
    fixed_table(doc, ["Model", "Trainable", "Model-only median (p95), ms", "End-to-end median (p95), ms", "End-to-end clips/s"], rows, [1.3, 0.9, 1.55, 1.65, 1.1])
    add_table_text(doc, "*Note.* Model-only timing starts from a device-resident normalized tensor. End-to-end timing includes original-video opening, decoding, frame sampling, letterboxing, normalization, transfer, forward pass, softmax, and explicit synchronization. This is offline clip latency, not a streaming benchmark.", above=False)


def add_reference(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    add_inline_runs(paragraph, text)


def build_body(
    source_path: Path,
    output_path: Path,
    figures_dir: Path,
    tokens: dict[str, Any],
    data: dict[str, pd.DataFrame],
    audits: dict[str, Any],
    replacements: dict[str, str],
    forbidden_identity: list[str],
) -> tuple[str, str, str, int]:
    source = replace_tokens(source_path.read_text(encoding="utf-8"), replacements)
    title, abstract, keywords, blocks = parse_blocks(source)
    keyword_values = [value.strip() for value in keywords.split(";") if value.strip()]
    if not 5 <= len(keyword_values) <= 7:
        raise ValueError(f"NCCSRJ requires 5–7 keywords; found {len(keyword_values)}")
    word_count = manuscript_word_count(abstract, blocks)
    if not 3000 <= word_count <= 5000:
        raise ValueError(f"NCCSRJ main-text word count must be 3,000–5,000; found {word_count}")

    doc = Document()
    configure_document(doc, tokens)
    set_core_properties(doc, title=title, author="", keywords=keywords)
    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_after = Pt(12)
    title_paragraph.paragraph_format.keep_with_next = True
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, size=12, bold=True)

    abstract_heading = doc.add_paragraph()
    abstract_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    abstract_heading.paragraph_format.keep_with_next = True
    abstract_heading.paragraph_format.line_spacing = 2.0
    set_run_font(abstract_heading.add_run("Abstract"), size=12, bold=True)
    add_body_paragraph(doc, abstract, first_line_indent=False)
    keyword_paragraph = add_body_paragraph(doc, "", first_line_indent=False)
    set_run_font(keyword_paragraph.add_run("Keywords: "), size=12, bold=True)
    set_run_font(keyword_paragraph.add_run(keywords), size=12)

    table_markers: dict[str, Callable[[], None]] = {
        "TABLE_PROTOCOL": lambda: table_protocol(doc, audits),
        "TABLE_MODELS": lambda: table_models(doc, data, audits),
        "TABLE_PRIMARY": lambda: table_primary(doc, data),
        "TABLE_FOLDS": lambda: table_folds(doc, data),
        "TABLE_CONDITIONS": lambda: table_conditions(doc, data),
        "TABLE_RUNTIME": lambda: table_runtime(doc, data),
    }
    figure_markers: dict[str, Callable[[], None]] = {
        "FIGURE_STUDY_FLOW": lambda: add_figure(
            doc,
            figures_dir / "figure_1_study_flow.png",
            "**Figure 1.** Audited NSL23 source flow and twofold signer-held-out design.",
            "Flow diagram from 630 NSL23 source videos to 375 eligible original isolated clips, followed by S1- and S2-held-out folds and 18 model runs.",
        ),
        "FIGURE_ARCHITECTURES": lambda: add_figure(
            doc,
            figures_dir / "figure_2_architectures.png",
            "**Figure 2.** Shared frozen MobileNetV2 encoder and the three evaluated classification heads.",
            "Architecture diagram comparing the middle-frame CNN, CNN–BiLSTM, and CNN–Transformer heads after a shared frozen MobileNetV2 encoder.",
        ),
        "FIGURE_PRIMARY": lambda: add_figure(
            doc,
            figures_dir / "figure_3_primary_performance.png",
            "**Figure 3.** Pooled out-of-fold top-1 accuracy and macro-F1 by architecture.",
            "Bar plots of accuracy and macro-F1 for three architectures, with seed-level points and standard-deviation error bars.",
        ),
        "FIGURE_TRAINING": lambda: add_figure(
            doc,
            figures_dir / "figure_4_training_curves.png",
            "**Figure 4.** Training loss, accuracy, and macro-F1 over the fixed 30-epoch schedule.",
            "Line plots of mean training loss, training accuracy, and training macro-F1 across folds and seeds for all three architectures, with standard-deviation bands.",
        ),
        "FIGURE_LIGHTING": lambda: add_figure(
            doc,
            figures_dir / "figure_5_lighting_performance.png",
            "**Figure 5.** Descriptive bright- and dark-condition performance.",
            "Grouped bars comparing bright and dark accuracy and macro-F1 across the three evaluated architectures.",
        ),
    }

    in_references = False
    reference_count = 0
    for kind, text in blocks:
        if kind == "heading1":
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline_runs(paragraph, text)
            in_references = text.lower() == "references"
        elif kind == "heading2":
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline_runs(paragraph, text)
        elif kind == "marker":
            if text in table_markers:
                table_markers[text]()
            elif text in figure_markers:
                figure_markers[text]()
            else:
                raise ValueError(f"Unknown manuscript marker: {text}")
        elif in_references:
            add_reference(doc, text)
            reference_count += 1
        else:
            add_body_paragraph(doc, text)
    if reference_count < 12:
        raise ValueError(f"Expected at least 12 APA references; found {reference_count}")

    searchable = source.lower()
    for value in forbidden_identity:
        if value and value.lower() in searchable:
            raise ValueError(f"Identity-sensitive content found in anonymous manuscript source: {value}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return title, abstract, keywords, word_count


def build_cover(
    output_path: Path,
    title: str,
    abstract: str,
    keywords: str,
    word_count: int,
    author: str,
    affiliation: str,
    email: str,
    tokens: dict[str, Any],
) -> None:
    if not all(value.strip() for value in (author, affiliation, email)):
        raise ValueError("Author, affiliation, and email are required for the separate cover sheet")
    doc = Document()
    configure_document(doc, tokens)
    set_core_properties(doc, title=title, author=author, keywords=keywords)

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_after = Pt(12)
    set_run_font(title_paragraph.add_run(title), size=12, bold=True)

    for value, bold in ((author, True), (affiliation, False), (f"Corresponding author: {email}", False)):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 2.0
        set_run_font(paragraph.add_run(value), size=12, bold=bold)

    abstract_heading = doc.add_paragraph()
    abstract_heading.paragraph_format.space_before = Pt(12)
    abstract_heading.paragraph_format.keep_with_next = True
    set_run_font(abstract_heading.add_run("Abstract"), size=12, bold=True)
    add_body_paragraph(doc, abstract, first_line_indent=False)
    keyword_paragraph = add_body_paragraph(doc, "", first_line_indent=False)
    set_run_font(keyword_paragraph.add_run("Keywords: "), size=12, bold=True)
    set_run_font(keyword_paragraph.add_run(keywords), size=12)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path("manuscript/blinded_manuscript.md"))
    parser.add_argument("--style-tokens", type=Path, default=Path("manuscript/nccsrj_style_tokens.json"))
    parser.add_argument("--results-dir", type=Path, default=Path("artifacts/results"))
    parser.add_argument("--runtime", type=Path, default=Path("artifacts/runtime/runtime_summary.csv"))
    parser.add_argument("--figures-dir", type=Path, default=Path("artifacts/figures"))
    parser.add_argument("--audit-dir", type=Path, default=Path("artifacts/audit"))
    parser.add_argument("--experiment-config", type=Path, default=Path("configs/experiment_local_m4.json"))
    parser.add_argument("--body-output", type=Path, default=Path("deliverables/NSL23_NCCSRJ_Blinded_Manuscript.docx"))
    parser.add_argument("--cover-output", type=Path, default=Path("deliverables/NSL23_NCCSRJ_Cover_Page.docx"))
    parser.add_argument("--author", required=True)
    parser.add_argument("--affiliation", required=True)
    parser.add_argument("--email", required=True)
    args = parser.parse_args()

    tokens = json.loads(args.style_tokens.read_text(encoding="utf-8"))
    data = load_results(args.results_dir, args.runtime)
    audits = load_audits(args.audit_dir, args.experiment_config)
    replacements = build_replacements(data, args.results_dir)
    title, abstract, keywords, word_count = build_body(
        args.source,
        args.body_output,
        args.figures_dir,
        tokens,
        data,
        audits,
        replacements,
        [args.author, args.affiliation, args.email, "/Users/"],
    )
    build_cover(
        args.cover_output,
        title,
        abstract,
        keywords,
        word_count,
        args.author,
        args.affiliation,
        args.email,
        tokens,
    )
    print(
        json.dumps(
            {
                "status": "pass",
                "body": str(args.body_output),
                "cover": str(args.cover_output),
                "main_text_words": word_count,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
